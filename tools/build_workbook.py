# -*- coding: utf-8 -*-
"""Build the monochrome Meiryo UI workbook from the companion Markdown."""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

ROOT = Path(__file__).resolve().parents[1]
MD_PATH = ROOT / "CloudWatchLogsからJP1までのアラート通知.md"
XLSX_PATH = ROOT / "CloudWatchLogsからJP1までのアラート通知.xlsx"

INK = "1A1A1A"
CHARCOAL = "2C2C2C"
SLATE = "3F3F3F"
MID = "5C5C5C"
LINE = "C6C6C2"
PAPER = "F7F7F5"
WHITE = "FFFFFF"
BAND = "ECECEA"
ZEBRA = "F3F3F1"
CODE = "EFEFED"

THIN = Border(
    left=Side(style="thin", color=LINE),
    right=Side(style="thin", color=LINE),
    top=Side(style="thin", color=LINE),
    bottom=Side(style="thin", color=LINE),
)
HAIR_BOTTOM = Border(bottom=Side(style="thin", color=LINE))


def font(size=10, bold=False, color=INK, underline=None):
    return Font(
        name="Meiryo UI",
        size=size,
        bold=bold,
        color=color,
        underline=underline,
        scheme=None,
    )


def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)


def align(horizontal="left", indent=0):
    return Alignment(
        horizontal=horizontal,
        vertical="center",
        wrap_text=True,
        indent=indent,
    )


def weight(text):
    total = 0
    for ch in str(text):
        total += 2 if ord(ch) > 127 else 1
    return total


def lines_for(text, excel_width):
    # 0.78 leaves padding so Meiryo UI full-width characters wrap inside the cell.
    capacity = max(6, int(float(excel_width) * 0.78))
    total = 0
    for para in str(text).split("\n"):
        if para == "":
            total += 1
            continue
        units = weight(para)
        total += max(1, (units + capacity - 1) // capacity)
    return max(1, total)


def height_for(text, excel_width, size=10, minimum=22, cap=180):
    lines = lines_for(text, excel_width)
    return min(cap, max(minimum, 8 + lines * (size + 7)))


def apply_page(ws, title):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.horizontalCentered = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(
        left=0.5, right=0.5, top=0.75, bottom=0.6, header=0.28, footer=0.28
    )
    ws.sheet_properties.tabColor = MID
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 110
    ws.sheet_view.view = "normal"
    ws.page_setup.scale = 70
    ws.oddHeader.left.text = '&"Meiryo UI"&9CloudWatch Logs から JP1 へのアラート'
    ws.oddHeader.right.text = '&"Meiryo UI"&9確認日 2026-10-07'
    ws.oddFooter.left.text = (
        '&"Meiryo UI"&8値は公式ドキュメントの確認日時点。Service Quotas と JP1 の版を正とする。'
    )
    ws.oddFooter.right.text = '&"Meiryo UI"&9&P / &N'
    ws.oddFooter.center.text = f'&"Meiryo UI"&9{title}'
    ws.sheet_format.defaultRowHeight = 18
    ws.print_options.horizontalCentered = True
    ws.page_setup.horizontalDpi = 300
    ws.page_setup.verticalDpi = 300
    ws.sheet_properties.outlinePr.summaryBelow = True


def paint(cell, value, *, size=10, bold=False, color=INK, bg=None, horizontal="left", border=True):
    if value not in (None, ""):
        cell.value = value
    cell.font = font(size, bold, color)
    cell.alignment = align(horizontal)
    if bg:
        cell.fill = fill(bg)
    if border:
        cell.border = THIN
    return cell


def paint_rest(ws, row, start_col, end_col, *, size=10, bold=False, color=INK, bg=None):
    for col in range(start_col, end_col + 1):
        cell = ws.cell(row, col)
        cell.font = font(size, bold, color)
        cell.alignment = align("left")
        if bg:
            cell.fill = fill(bg)
        cell.border = THIN


def set_widths(ws, widths):
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    last = get_column_letter(len(widths))
    for col in range(len(widths) + 1, 12):
        ws.column_dimensions[get_column_letter(col)].width = 3
    ws.auto_filter.ref = None
    ws.page_setup.fitToWidth = 1
    ws.print_title_rows = "1:2"
    ws.freeze_panes = "A3"
    ws.sheet_view.zoomScale = 110
    return last


def merge_write(ws, row, cols, value, *, size=10, bold=False, color=INK, bg=None, height=None, min_height=20):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    cell = ws.cell(row, 1)
    paint(cell, value, size=size, bold=bold, color=color, bg=bg, border=bg is not None)
    if bg is None:
        cell.border = HAIR_BOTTOM
    total = sum(ws.column_dimensions[get_column_letter(i)].width or 12 for i in range(1, cols + 1))
    ws.row_dimensions[row].height = height or height_for(value, total, size, min_height)
    edge = THIN if bg else HAIR_BOTTOM
    for col in range(2, cols + 1):
        other = ws.cell(row, col)
        other.border = edge
        if bg:
            other.fill = fill(bg)
    return row + 1


def write_table(ws, row, rows, widths):
    header = rows[0]
    cols = len(header)
    for col, text in enumerate(header, start=1):
        paint(
            ws.cell(row, col),
            text,
            size=10,
            bold=True,
            color=WHITE,
            bg=CHARCOAL,
            horizontal="center",
        )
    ws.row_dimensions[row].height = height_for("\n".join(header), min(widths), 10, 24)
    row += 1
    first_data = row
    for offset, data in enumerate(rows[1:]):
        bg = WHITE if offset % 2 == 0 else ZEBRA
        padded = list(data) + [""] * (cols - len(data))
        max_h = 20
        for col, text in enumerate(padded, start=1):
            paint(ws.cell(row, col), text, size=10, bg=bg)
            max_h = max(max_h, height_for(text, widths[col - 1], 10, 20))
        ws.row_dimensions[row].height = min(120, max_h)
        row += 1
    last_data = row - 1
    if last_data >= first_data:
        ws.auto_filter.ref = f"A{first_data - 1}:{get_column_letter(cols)}{last_data}"
        ws.auto_filter.ref = f"A{first_data - 1}:{get_column_letter(cols)}{last_data}"
    return row


def strip_md(text):
    text = text.replace("**", "")
    text = text.replace("`", "")
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1（\2）", text)
    return text.strip()


def parse_table(lines):
    parsed = []
    for line in lines:
        if re.match(r"^\s*\|?\s*:?-{3,}", line.replace("|", " | ")):
            continue
        if set(line.replace("|", "").replace(":", "").replace("-", "").strip()) == set():
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) or c == "" for c in cells):
                continue
        cells = [strip_md(c.strip()) for c in line.strip().strip("|").split("|")]
        if cells and all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c != ""):
            continue
        parsed.append(cells)
    return parsed


def parse_sections(markdown):
    lines = markdown.splitlines()
    preamble = []
    sections = []
    current = None
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("## "):
            current = {"title": line[3:].strip(), "blocks": []}
            sections.append(current)
            i += 1
            continue
        if current is None:
            preamble.append(line)
            i += 1
            continue
        if line.startswith("### "):
            current["blocks"].append(("h3", line[4:].strip()))
            i += 1
            continue
        if line.startswith("```"):
            buff = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                buff.append(lines[i])
                i += 1
            i += 1
            current["blocks"].append(("code", "\n".join(buff)))
            continue
        if line.strip().startswith("|"):
            buff = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                buff.append(lines[i])
                i += 1
            current["blocks"].append(("table", parse_table(buff)))
            continue
        if re.match(r"^(\d+\.\s+|- )", line.strip()):
            buff = [line.strip()]
            i += 1
            while i < len(lines) and re.match(r"^(\d+\.\s+|- )", lines[i].strip()):
                buff.append(lines[i].strip())
                i += 1
            current["blocks"].append(("list", buff))
            continue
        if line.strip() in ("", "---", "***", "___"):
            i += 1
            continue
        buff = [line.strip()]
        i += 1
        while (
            i < len(lines)
            and lines[i].strip()
            and lines[i].strip() not in ("---", "***", "___")
            and not lines[i].startswith("#")
            and not lines[i].strip().startswith("|")
            and not lines[i].startswith("```")
            and not re.match(r"^(\d+\.\s+|- )", lines[i].strip())
        ):
            buff.append(lines[i].strip())
            i += 1
        current["blocks"].append(("p", " ".join(buff)))
    return "\n".join(preamble), sections


SHEET_NAMES = [
    "01_たとえ話",
    "02_歴史と背景",
    "03_情報の変身",
    "04_具体例",
    "05_時刻の進み方",
    "06_CloudWatchLogs",
    "07_メトリクスフィルタ",
    "08_パターン",
    "09_アラーム",
    "10_ログアラーム",
    "11_SNS",
    "12_アカウントをまたぐ",
    "13_SQS",
    "14_EC2の仕事",
    "15_JP1",
    "16_失敗と対策",
    "17_決めること",
    "18_出典",
]


def column_widths(section, cols):
    widths = [18] * cols
    samples = [[] for _ in range(cols)]
    for kind, payload in section["blocks"]:
        if kind != "table":
            continue
        for record in payload:
            for index, cell in enumerate(record[:cols]):
                samples[index].append(cell)
    for index in range(cols):
        longest = max((weight(text) for text in samples[index]), default=16)
        # Excel width is about one half-width character.
        widths[index] = min(46, max(16, int(longest / 2) + 3))
    # Keep the sheet printable. Shrink the widest columns together if needed.
    total = sum(widths)
    if total > 210:
        scale = 210 / total
        widths = [max(14, int(w * scale)) for w in widths]
    return widths


def render_section(wb, sheet_name, section):
    ws = wb.create_sheet(sheet_name)
    apply_page(ws, section["title"])
    tables = [block for kind, block in section["blocks"] if kind == "table"]
    cols = max((len(table[0]) for table in tables if table), default=1)
    cols = max(1, cols)
    widths = column_widths(section, cols) if tables else [128]
    if not tables:
        cols = 1
    set_widths(ws, widths)
    ws.auto_filter.ref = None
    row = 1
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=cols)
    paint(ws.cell(1, 1), section["title"], size=16, bold=True, color=WHITE, bg=INK)
    for col in range(2, cols + 1):
        paint(ws.cell(1, col), "", size=16, bold=True, color=WHITE, bg=INK)
    ws.row_dimensions[1].height = 30
    row = 2
    ws.row_dimensions[2].height = 8
    row = 3
    last_table_is_final = bool(section["blocks"]) and section["blocks"][-1][0] == "table"
    seen_tables = 0
    table_total = sum(1 for kind, _ in section["blocks"] if kind == "table")
    for kind, payload in section["blocks"]:
        if kind == "h3":
            row = merge_write(
                ws, row, cols, payload, size=12, bold=True, color=WHITE, bg=SLATE, min_height=24
            )
            continue
        if kind == "p":
            text = strip_md(payload)
            row = merge_write(ws, row, cols, text, size=10, bg=PAPER, min_height=22)
            continue
        if kind == "list":
            for item in payload:
                text = strip_md(item)
                if text.startswith("- "):
                    text = "・" + text[2:]
                row = merge_write(ws, row, cols, text, size=10, bg=WHITE, min_height=22)
            continue
        if kind == "code":
            for code_line in payload.splitlines() or [""]:
                row = merge_write(ws, row, cols, code_line if code_line else " ", size=9, bg=CODE, min_height=18)
            continue
        if kind == "table":
            seen_tables += 1
            header = payload[0]
            data_rows = payload[1:]
            # Pad ragged rows.
            width = len(header)
            normalized = [header]
            for record in data_rows:
                normalized.append((record + [""] * width)[:width])
            start = row
            for col, text in enumerate(header, start=1):
                paint(ws.cell(row, col), text, size=10, bold=True, color=WHITE, bg=CHARCOAL, horizontal="center")
            ws.row_dimensions[row].height = max(22, height_for(" / ".join(header), 24, 10, 22))
            row += 1
            first = row
            for offset, record in enumerate(normalized[1:]):
                bg = WHITE if offset % 2 == 0 else ZEBRA
                max_h = 20
                for col, text in enumerate(record, start=1):
                    paint(ws.cell(row, col), text, size=10, bg=bg)
                    max_h = max(max_h, height_for(text, widths[col - 1] if col - 1 < len(widths) else 20, 10, 20))
                ws.row_dimensions[row].height = min(110, max_h)
                row += 1
            last = row - 1
            # Only the final table on a sheet gets a filter, so filtering cannot hide later paragraphs.
            if last_table_is_final and seen_tables == table_total and last >= first:
                ws.auto_filter.ref = f"A{start}:{get_column_letter(width)}{last}"
            row += 1
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = ws.auto_filter.ref
    ws.oddHeader.center.text = f'&"Meiryo UI"&9{sheet_name}'
    ws.page_setup.fitToWidth = 1
    return ws


def build_cover(wb, intro):
    ws = wb.active
    ws.title = "00_表紙"
    apply_page(ws, "表紙")
    widths = [28, 36, 36, 36, 28]
    set_widths(ws, widths)
    ws.freeze_panes = "A8"
    ws.print_title_rows = "1:1"
    ws.merge_cells("A1:E1")
    paint(ws.cell(1, 1), "CloudWatch Logs から JP1 までのアラート通知", size=20, bold=True, color=WHITE, bg=INK)
    for col in range(2, 6):
        paint(ws.cell(1, col), "", bg=INK, color=WHITE)
    ws.row_dimensions[1].height = 40
    ws.merge_cells("A2:E2")
    paint(
        ws.cell(2, 1),
        "ログの文字列を数え、閾値を超えたら、別アカウントの郵便受けを経由して JP1 の当番ノートへ届ける",
        size=12,
        bold=True,
        color=INK,
        bg=BAND,
    )
    for col in range(2, 6):
        paint(ws.cell(2, col), "", bg=BAND)
    ws.row_dimensions[2].height = 28
    meta = [
        ("確認日", "2026-10-07"),
        ("フォント", "Meiryo UI。色は白、灰、黒だけ"),
        ("対になる文書", "同じフォルダの Markdown"),
        ("リージョンの例", "ap-northeast-1（東京）"),
        ("アカウントの例", "ワークロード 111122223333 ／ 共通運用 444455556666"),
    ]
    for index, (label, value) in enumerate(meta, start=4):
        paint(ws.cell(index, 1), label, bold=True, color=WHITE, bg=SLATE, horizontal="center")
        ws.merge_cells(start_row=index, start_column=2, end_row=index, end_column=5)
        paint(ws.cell(index, 2), value, bg=WHITE)
        for col in range(3, 6):
            paint(ws.cell(index, col), "", bg=WHITE)
        ws.row_dimensions[index].height = 22
    ws.merge_cells("A10:E10")
    paint(
        ws.cell(10, 1),
        "最初に覚えること",
        size=14,
        bold=True,
        color=WHITE,
        bg=CHARCOAL,
    )
    for col in range(2, 6):
        paint(ws.cell(10, col), "", bg=CHARCOAL)
    ws.row_dimensions[10].height = 26
    points = [
        "日記帳は CloudWatch Logs。文章のまま残る。",
        "数え係はメトリクスフィルタ。新しく入った行だけを見て、一致を点数にする。過去の日記は数え直さない。",
        "先生のルールは CloudWatch アラーム。点数が閾値を超えた状態へ変わった瞬間だけ、放送する。鳴り続けるベルではない。",
        "校内放送は SNS。職員室の郵便受けは、別の家（共通運用アカウント）にある SQS。",
        "連絡係の EC2 が、手紙を JP1 イベントの言葉へ翻訳する。当番ノートの入口は JP1/Base、一覧は JP1/IM。",
        "この経路の手紙に、ログの原文は原則として入らない。原文を当番ノートへ載せるには、説明文・EC2 の読み直し・ログアラームの行添付のどれかを選ぶ。",
    ]
    for index, text in enumerate(points):
        row = 11 + index
        paint(ws.cell(row, 1), str(index + 1), bold=True, color=WHITE, bg=MID, horizontal="center")
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
        paint(ws.cell(row, 2), text, bg=PAPER if index % 2 == 0 else WHITE)
        for col in range(3, 6):
            paint(ws.cell(row, col), "", bg=PAPER if index % 2 == 0 else WHITE)
        ws.row_dimensions[row].height = height_for(text, 136, 10, 24)
    ws.merge_cells("A18:E18")
    paint(ws.cell(18, 1), "色の意味", size=12, bold=True, color=WHITE, bg=CHARCOAL)
    for col in range(2, 6):
        paint(ws.cell(18, col), "", bg=CHARCOAL)
    legend = [
        ("黒", "章のタイトル"),
        ("濃い灰", "表の見出し、または小見出し"),
        ("薄い灰", "説明、または一行おきの表"),
        ("白", "値、手順、本文"),
        ("枠線", "項目の区切り。フィルタを付けた表は、その表がシートの最後にあるときだけ"),
    ]
    for index, (name, meaning) in enumerate(legend):
        row = 19 + index
        paint(ws.cell(row, 1), name, bold=True, horizontal="center", bg=BAND)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
        paint(ws.cell(row, 2), meaning, bg=WHITE)
        for col in range(3, 6):
            paint(ws.cell(row, col), "", bg=WHITE)
        ws.row_dimensions[row].height = 22
    ws.merge_cells("A25:E25")
    paint(
        ws.cell(25, 1),
        "各シートに本文と表がある。表がそのシートの最後にあるときだけ、見出しで絞り込みができる。値は確認日時点のもので、Service Quotas と JP1 の版が優先する。",
        size=10,
        bg=PAPER,
    )
    for col in range(2, 6):
        paint(ws.cell(25, col), "", bg=PAPER)
    ws.row_dimensions[25].height = height_for(
        "各シートに本文と表がある。表がそのシートの最後にあるときだけ、見出しで絞り込みができる。値は確認日時点のもので、Service Quotas と JP1 の版が優先する。",
        136,
        10,
        28,
    )
    ws.row_dimensions[3].height = 10


def build_toc(wb, names):
    ws = wb.create_sheet("00b_目次", 1)
    apply_page(ws, "目次")
    widths = [16, 42, 78]
    set_widths(ws, widths)
    headers = ["シート", "見出し", "このシートで分かること"]
    blurbs = [
        "学校の日記、数え係、放送、郵便受け、当番ノートへの対応",
        "JP1、SQS、CloudWatch、SNS、Logs、マルチアカウントが必要になった順番",
        "段ごとに何が残り、何が点数へ消えるか",
        "5分で ERROR が6回のとき、JSON と jevsend に何が乗るか",
        "1分の箱、評価、SQS の待ち。秒では検知しない理由",
        "ロググループ、保持、イベントサイズ、クラス",
        "フィルタの各フィールド、デフォルト値、ディメンション、Sum",
        "語、JSON、スペース区切り、正規表現の書き方と上限",
        "状態、M out of N、期間、アクションは遷移のときだけ",
        "クエリで判定し、最大50行を通知へ添付する新しいベル",
        "標準トピック、件名100文字、Publish の許可、MaximumMessageSize",
        "トピック、購読、キューの三つの許可と KMS",
        "可視性、保持、ロングポーリング、DLQ",
        "翻訳の順番、対応表、原文を読み直す場合",
        "jevsend、重大度、1,023バイト、20098/tcp",
        "黙る原因と、先に置くもの",
        "実装に渡す前に決める項目と、起点の値",
        "確認に使った公式ドキュメント",
    ]
    paint(ws.cell(1, 1), "目次", size=16, bold=True, color=WHITE, bg=INK)
    paint(ws.cell(1, 2), "", bg=INK)
    paint(ws.cell(1, 3), "シート名をクリックするとその章へ移動する", size=11, bold=True, color=WHITE, bg=INK)
    ws.row_dimensions[1].height = 30
    for col, text in enumerate(headers, start=1):
        paint(ws.cell(3, col), text, bold=True, color=WHITE, bg=CHARCOAL, horizontal="center")
    ws.row_dimensions[3].height = 22
    for index, name in enumerate(names):
        row = 4 + index
        bg = WHITE if index % 2 == 0 else ZEBRA
        paint(ws.cell(row, 1), f"{index + 1:02d}", bold=True, bg=bg, horizontal="center")
        link = ws.cell(row, 2)
        paint(link, name, bg=bg)
        link.hyperlink = f"#'{name}'!A1"
        link.font = font(10, color=INK, underline="single")
        paint(ws.cell(row, 3), blurbs[index], bg=bg)
        ws.row_dimensions[row].height = height_for(blurbs[index], widths[2], 10, 22)
    ws.auto_filter.ref = f"A3:C{3 + len(names)}"
    ws.freeze_panes = "A4"


def build_flow(wb):
    ws = wb.create_sheet("00c_全体の流れ", 2)
    apply_page(ws, "全体の流れ")
    labels = ["アプリ", "Logs", "フィルタ", "メトリクス", "アラーム", "SNS", "SQS", "EC2", "JP1/Base", "JP1/IM"]
    places = ["ワークロードアカウント", "", "", "", "", "同じアカウント\n同じリージョン", "共通運用\nアカウント", "共通運用\nアカウント", "通知サーバ\nまたは IM", "IM サーバ"]
    keeps = ["文章", "原文を保存", "一致を点数化", "期間の合計", "状態の変化", "JSON を複製", "手紙を保管", "JP1 の言葉へ", "イベント登録", "画面と対処"]
    drops = ["", "", "原文は残さない", "どの行かは残らない", "ロググループ名は説明が無いと消える", "ちょうど1回ではない", "順序は保証されない", "1,023バイト超", "転送条件の外", "上がらなかったもの"]
    widths = [16] * 10
    set_widths(ws, widths)
    ws.merge_cells("A1:J1")
    paint(ws.cell(1, 1), "左から右へ、一通のアラートが進む", size=16, bold=True, color=WHITE, bg=INK)
    for col in range(2, 11):
        paint(ws.cell(1, col), "", bg=INK)
    ws.row_dimensions[1].height = 30
    ws.merge_cells("A2:J2")
    paint(
        ws.cell(2, 1),
        "黒い箱は文章を扱う段。灰色の箱は点数と通知の段。濃い灰色は共通運用アカウント以降。",
        size=10,
        bg=BAND,
    )
    for col in range(2, 11):
        paint(ws.cell(2, col), "", bg=BAND)
    ws.row_dimensions[2].height = 24
    tones = [INK, INK, SLATE, SLATE, SLATE, SLATE, CHARCOAL, CHARCOAL, CHARCOAL, CHARCOAL]
    row_labels = [("段", labels, True), ("いる場所", places, False), ("この段が残すもの", keeps, False), ("この段で落ちるもの", drops, False)]
    start = 4
    for offset, (caption, values, dark_text_white) in enumerate(row_labels):
        row = start + offset
        max_h = 32
        for col, text in enumerate(values, start=1):
            bg = tones[col - 1] if offset == 0 else (WHITE if offset % 2 else BAND)
            color = WHITE if offset == 0 else INK
            paint(ws.cell(row, col), text, size=10, bold=(offset == 0), color=color, bg=bg, horizontal="center")
            max_h = max(max_h, height_for(text or " ", widths[col - 1], 10, 32))
        ws.row_dimensions[row].height = min(78, max_h)
    notes = [
        "フィルタの点数は1分の箱に入る。アラームの期間は 60 の倍数秒にする。10秒、20秒、30秒は高解像度メトリクス用で、この点数には使わない。",
        "アラームのアクションは、OK、ALARM、INSUFFICIENT_DATA のあいだを移ったときだけ動く。ALARM のまま件数が多い状態が続いても、二通目は出ない。",
        "SNS は標準トピック。FIFO は使わない。SQS も標準。購読はキューの持ち主である共通運用アカウントが作ると、確認が自動で終わる。",
        "EC2 は jevsendd の成功を確認してから DeleteMessage する。失敗した手紙は可視性タイムアウトのあと戻り、規定回数でデッドレターへ移る。",
    ]
    for index, text in enumerate(notes):
        row = 9 + index
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
        paint(ws.cell(row, 1), text, size=10, bg=PAPER)
        for col in range(2, 11):
            paint(ws.cell(row, col), "", bg=PAPER)
        ws.row_dimensions[row].height = height_for(text, 150, 10, 28)
    ws.freeze_panes = "A4"


def build_baseline(wb):
    ws = wb.create_sheet("19_起点の値")
    apply_page(ws, "起点の値")
    headers = ["場所", "項目", "起点として置く値", "そうする理由"]
    rows = [
        ["フィルタ", "パターン", "アプリが実際に出す語。例: ERROR", "大文字と小文字は区別される"],
        ["フィルタ", "metricValue", "1", "1件の一致を1点にする"],
        ["フィルタ", "defaultValue", "0", "ログは来たが一致が無い分を0にする。ディメンションがあるときは付けられない"],
        ["フィルタ", "dimensions", "なし", "種類が増えると別メトリクスになり、1時間に1000組でフィルタが止まり得る"],
        ["フィルタ", "unit", "Count", "あとから変えても効かない。アラームと揃える"],
        ["アラーム", "Statistic", "Sum", "期間中の件数になる。Average は件数にならない"],
        ["アラーム", "Period", "300", "5分合計。60の倍数だけを使う"],
        ["アラーム", "比較と閾値", "GreaterThanThreshold と、当番が動く件数", "少なくとも1回の二重計上があるので、ぴったり一致では監査しない"],
        ["アラーム", "M / N", "2 / 3", "直近3枠のうち2枠。連続でなくてよい"],
        ["アラーム", "TreatMissingData", "notBreaching", "静かな時間を正常とみなす"],
        ["アラーム", "OKActions", "復旧を出すなら AlarmActions と同じトピック", "付けないと復旧は JP1 に届かない"],
        ["アラーム", "InsufficientDataActions", "空", "作成直後のデータ不足で鳴らせない"],
        ["監視の穴", "IncomingLogEvents", "そのロググループで別アラーム", "ログ停止をエラー件数の OK と区別する"],
        ["SNS", "種類", "標準トピック。アラームと同じリージョン、同じアカウント", "FIFO はアラームの発行に合わない"],
        ["SQS", "保持", "14日", "既定の4日は連休を超えると消える"],
        ["SQS", "受信待ち", "20秒", "ロングポーリング"],
        ["SQS", "可視性", "120秒から測って調整", "jevsendd が終わるより長く。最大12時間"],
        ["SQS", "DLQ", "maxReceiveCount 5", "壊れた手紙で本線を埋めない"],
        ["EC2", "削除", "jevsendd の成功後", "失敗はキューへ戻す"],
        ["JP1", "重大度", "障害は Error。復旧は転送される値", "既定では Notice と Information はマネージャーへ上がらない"],
        ["JP1", "メッセージ", "対応表の日本語。1,023バイトで切る", "文字数ではなく、イベントの文字コードのバイト数"],
        ["ネットワーク", "ポート", "通知サーバから IM の 20098/tcp のみ", "イベント連携。0.0.0.0/0 には開けない"],
    ]
    widths = [18, 28, 62, 62]
    set_widths(ws, widths)
    ws.merge_cells("A1:D1")
    paint(ws.cell(1, 1), "実装に入るとき、まずこの値で置き、要件で数字だけ変える", size=16, bold=True, color=WHITE, bg=INK)
    for col in range(2, 5):
        paint(ws.cell(1, col), "", bg=INK)
    ws.row_dimensions[1].height = 30
    for col, text in enumerate(headers, start=1):
        paint(ws.cell(3, col), text, bold=True, color=WHITE, bg=CHARCOAL, horizontal="center")
    ws.row_dimensions[3].height = 22
    for offset, record in enumerate(rows):
        row = 4 + offset
        bg = WHITE if offset % 2 == 0 else ZEBRA
        max_h = 20
        for col, text in enumerate(record, start=1):
            paint(ws.cell(row, col), text, bg=bg, horizontal="center" if col == 1 else "left")
            max_h = max(max_h, height_for(text, widths[col - 1], 10, 20))
        ws.row_dimensions[row].height = min(72, max_h)
    ws.auto_filter.ref = f"A3:D{3 + len(rows)}"
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:D{3 + len(rows)}"


def assert_workbook(wb):
    offenses = []
    required = [
        "メトリクスフィルタ",
        "notBreaching",
        "20098",
        "1,023",
        "ActionLogLineCount",
        "jevsendd",
        "MaximumMessageSize",
        "IncomingLogEvents",
    ]
    blob = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.font and cell.font.name and cell.font.name != "Meiryo UI":
                    offenses.append(f"font {cell.font.name} at {ws.title}!{cell.coordinate}")
                if cell.font and cell.font.scheme not in (None, "none"):
                    offenses.append(f"scheme {cell.font.scheme} at {ws.title}!{cell.coordinate}")
                if cell.fill and cell.fill.fgColor is not None and cell.fill.patternType == "solid":
                    rgb = cell.fill.fgColor.rgb
                    if rgb and isinstance(rgb, str) and len(rgb) >= 6:
                        hex_part = rgb[-6:]
                        if len(hex_part) == 6 and all(c in "0123456789ABCDEFabcdef" for c in hex_part):
                            r = int(hex_part[0:2], 16)
                            g = int(hex_part[2:4], 16)
                            b = int(hex_part[4:6], 16)
                            if abs(r - g) > 2 or abs(g - b) > 2:
                                offenses.append(f"color {hex_part} at {ws.title}!{cell.coordinate}")
                if cell.value is not None:
                    blob.append(str(cell.value))
                if cell.alignment and cell.alignment.wrap_text is not True and cell.value not in (None, ""):
                    offenses.append(f"nowrap {ws.title}!{cell.coordinate}")
                if isinstance(cell.value, str) and cell.value.strip() in ("---", "***"):
                    offenses.append(f"rule left at {ws.title}!{cell.coordinate}")
                if isinstance(cell.value, str) and len(cell.value) > 40:
                    span = None
                    for merged in ws.merged_cells.ranges:
                        if cell.coordinate in merged:
                            if cell.row != merged.min_row or cell.column != merged.min_col:
                                span = "skip"
                            else:
                                span = sum(
                                    ws.column_dimensions[get_column_letter(col)].width or 10
                                    for col in range(merged.min_col, merged.max_col + 1)
                                )
                            break
                    if span == "skip":
                        continue
                    excel_width = span or (ws.column_dimensions[cell.column_letter].width or 10)
                    needed = height_for(cell.value, excel_width, cell.font.size or 10, 18, cap=1000)
                    actual = ws.row_dimensions[cell.row].height or 15
                    if actual + 1 < needed * 0.92:
                        offenses.append(
                            f"clip {ws.title}!{cell.coordinate} h={actual:.0f} need={needed:.0f}"
                        )
    text = "\n".join(blob)
    for token in required:
        if token not in text:
            offenses.append(f"missing token {token}")
    if offenses:
        raise SystemExit("workbook check failed:\n" + "\n".join(offenses[:40]))


def main():
    markdown = MD_PATH.read_text(encoding="utf-8")
    intro, sections = parse_sections(markdown)
    if len(sections) != len(SHEET_NAMES):
        raise SystemExit(f"section count {len(sections)} != sheet names {len(SHEET_NAMES)}")
    wb = Workbook()
    default_font = wb._fonts[0]
    default_font.name = "Meiryo UI"
    default_font.sz = 10
    default_font.scheme = None
    build_cover(wb, intro)
    build_toc(wb, SHEET_NAMES)
    build_flow(wb)
    for name, section in zip(SHEET_NAMES, sections):
        render_section(wb, name, section)
    build_baseline(wb)
    # Re-apply default font after sheets exist.
    wb._fonts[0].name = "Meiryo UI"
    wb._fonts[0].sz = 10
    wb._fonts[0].scheme = None
    assert_workbook(wb)
    wb.properties.title = "CloudWatch Logs から JP1 までのアラート通知"
    wb.properties.creator = "Grok"
    wb.properties.subject = "メトリクスフィルタ、アラーム、SNS、クロスアカウント SQS、JP1"
    wb.properties.description = "確認日 2026-10-07。Meiryo UI、モノトーン。"
    wb.properties.keywords = "CloudWatch,SNS,SQS,JP1,メトリクスフィルタ"
    wb.properties.category = "運用設計"
    XLSX_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(XLSX_PATH)
    print(f"sheets={len(wb.worksheets)}")
    print(f"wrote={XLSX_PATH}")


if __name__ == "__main__":
    main()
